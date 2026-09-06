#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
N3 読解 Excel → Google Sheets インポート

使い方:
  pip install openpyxl google-auth google-auth-oauthlib google-api-python-client
  python import_reading_to_sheets.py --dry-run
  python import_reading_to_sheets.py --create
  python import_reading_to_sheets.py
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
CONFIG_FILE = SCRIPT_DIR / "reading_config.json"
SCOPES = ["https://www.googleapis.com/auth/spreadsheets"]

PRACTICE_HEADERS = [
    "週", "日目", "テーマ", "テーマ（英）", "問題種別", "本文", "問番号",
    "設問文", "設問文（英）", "選択肢1", "選択肢2", "選択肢3", "選択肢4", "解答",
]
EXAM_HEADERS = [
    "週", "問番号", "本文", "設問文", "設問文（英）",
    "選択肢1", "選択肢2", "選択肢3", "選択肢4", "解答",
]

THEME_EN = {
    "案内①": "Notice 1", "案内②": "Notice 2", "案内③": "Notice 3",
    "試験要項": "Exam guidelines", "募集①": "Recruitment 1", "募集②": "Recruitment 2",
    "カタログ①": "Catalog 1", "カタログ②": "Catalog 2",
    "グラフ①": "Graph 1", "グラフ②": "Graph 2",
    "メール①": "Email 1", "メール②": "Email 2",
    "家族①": "Family 1", "家族②": "Family 2",
    "小説①": "Novel 1", "小説②": "Novel 2",
    "広告①": "Advertisement 1", "広告②": "Advertisement 2",
    "意見文①": "Opinion essay 1", "意見文②": "Opinion essay 2", "意見文③": "Opinion essay 3",
    "手紙・はがき①": "Letter/Postcard 1", "手紙・はがき②": "Letter/Postcard 2", "手紙・はがき③": "Letter/Postcard 3",
    "日記①": "Diary 1", "日記②": "Diary 2",
    "説明書①": "Manual 1", "説明書②": "Manual 2",
    "お知らせ": "Announcement", "まんが": "Manga", "保証書": "Warranty",
    "見出し": "Headline", "FAX（ビジネスレター）": "FAX (business letter)",
    "医学に関する文章": "Medical text", "社会に関する文章": "Social issues text",
    "計算に関する文章": "Calculation text",
}

QUESTION_EN = {
    "次の会話文を読んで、後の文から正しいものを選ぼう。": "Read the conversation and choose the correct statement.",
    "この案内の内容と合っているものはどれか。": "Choose the statement that matches the notice.",
}


def load_config() -> dict:
    with CONFIG_FILE.open(encoding="utf-8") as f:
        return json.load(f)


def save_config(cfg: dict) -> None:
    with CONFIG_FILE.open("w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


def cell(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def to_int_answer(value) -> str:
    if value is None or value == "":
        return ""
    try:
        return str(int(float(value)))
    except (TypeError, ValueError):
        return cell(value)


def get_sheets_service(credentials: str, token: str):
    from google.auth.transport.requests import Request
    from google.oauth2.credentials import Credentials
    from google_auth_oauthlib.flow import InstalledAppFlow
    from googleapiclient.discovery import build

    cred_path = Path(credentials)
    token_path = Path(token)
    creds = None
    if token_path.exists():
        creds = Credentials.from_authorized_user_file(str(token_path), SCOPES)
    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            if not cred_path.exists():
                raise FileNotFoundError(f"credentials.json が見つかりません: {cred_path}")
            flow = InstalledAppFlow.from_client_secrets_file(str(cred_path), SCOPES)
            creds = flow.run_local_server(port=0)
        token_path.write_text(creds.to_json(), encoding="utf-8")
    return build("sheets", "v4", credentials=creds)


def read_excel(path: str) -> tuple[list[list], list[list]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    practice_rows: list[list] = []
    exam_rows: list[list] = []

    ws = wb["読解問題"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[4]:
            continue
        theme = cell(row[2])
        question = cell(row[6])
        practice_rows.append([
            cell(row[0]), cell(row[1]), theme, THEME_EN.get(theme, ""),
            cell(row[3]), cell(row[4]), cell(row[5]), question,
            QUESTION_EN.get(question, ""),
            cell(row[7]), cell(row[8]), cell(row[9]), cell(row[10]),
            to_int_answer(row[12]),
        ])

    ws = wb["実戦問題"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[2]:
            continue
        question = cell(row[3])
        exam_rows.append([
            cell(row[0]), cell(row[1]), cell(row[2]), question,
            QUESTION_EN.get(question, ""),
            cell(row[4]), cell(row[5]), cell(row[6]), cell(row[7]),
            to_int_answer(row[8]),
        ])

    return practice_rows, exam_rows


def create_spreadsheet(service, title: str) -> str:
    body = {
        "properties": {"title": title},
        "sheets": [
            {"properties": {"title": "読解問題"}},
            {"properties": {"title": "実戦問題"}},
        ],
    }
    ss = service.spreadsheets().create(body=body, fields="spreadsheetId").execute()
    return ss["spreadsheetId"]


def write_sheet(service, spreadsheet_id: str, sheet_name: str, headers: list, rows: list[list]) -> None:
    values = [headers] + rows
    service.spreadsheets().values().clear(
        spreadsheetId=spreadsheet_id,
        range=f"'{sheet_name}'!A:Z",
    ).execute()
    service.spreadsheets().values().update(
        spreadsheetId=spreadsheet_id,
        range=f"'{sheet_name}'!A1",
        valueInputOption="RAW",
        body={"values": values},
    ).execute()


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true", help="読み込み件数のみ表示")
    parser.add_argument("--create", action="store_true", help="新規スプレッドシートを作成")
    args = parser.parse_args()

    cfg = load_config()
    excel_path = cfg.get("excel_path", "")
    practice_rows, exam_rows = read_excel(excel_path)

    print(f"読解問題: {len(practice_rows)} 行")
    print(f"実戦問題: {len(exam_rows)} 行")

    if args.dry_run:
        return

    service = get_sheets_service(cfg["credentials"], cfg["token"])
    spreadsheet_id = cfg.get("spreadsheet_id", "")

    if args.create or not spreadsheet_id:
        spreadsheet_id = create_spreadsheet(service, "N3 読解データ")
        cfg["spreadsheet_id"] = spreadsheet_id
        save_config(cfg)
        print(f"新規スプレッドシート作成: {spreadsheet_id}")

    write_sheet(service, spreadsheet_id, "読解問題", PRACTICE_HEADERS, practice_rows)
    write_sheet(service, spreadsheet_id, "実戦問題", EXAM_HEADERS, exam_rows)
    print(f"アップロード完了: https://docs.google.com/spreadsheets/d/{spreadsheet_id}/edit")
    print("Reading.gs の SPREADSHEET_IDS.reading に上記 ID を設定してください。")


if __name__ == "__main__":
    main()
